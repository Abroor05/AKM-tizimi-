import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_tz_document(output_path):
    doc = docx.Document()

    # Sahifa chetki masofalari (A4 standarti: chap 3 cm, o'ng 1.5 cm, yuqori 2 cm, past 2 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)     # 2.0 cm
        section.bottom_margin = Inches(0.79)  # 2.0 cm
        section.left_margin = Inches(1.18)    # 3.0 cm
        section.right_margin = Inches(0.59)   # 1.5 cm
        section.page_width = Inches(8.27)     # 21.0 cm (A4)
        section.page_height = Inches(11.69)   # 29.7 cm (A4)

    # Normal uslubni sozlash
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(3)
    normal_style.paragraph_format.space_before = Pt(0)

    def apply_run_font(run, font_name="Times New Roman", size_pt=14, bold=False, italic=False, color=RGBColor(0, 0, 0)):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
        rPr.append(rFonts)

    def add_p(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, indent=True, space_before=0, space_after=3):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.49) # 1.25 cm xatboshi
        else:
            p.paragraph_format.first_line_indent = Inches(0)
        run = p.add_run(text)
        apply_run_font(run, font_name="Times New Roman", size_pt=14, bold=bold, italic=italic)
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        apply_run_font(run, font_name="Times New Roman", size_pt=14, bold=True)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Inches(0.49)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        apply_run_font(run, font_name="Times New Roman", size_pt=14, bold=True)
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.first_line_indent = Inches(0.49)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        apply_run_font(run, font_name="Times New Roman", size_pt=14, bold=True, italic=True)
        return p

    # --- Sarlavha sahifasi / Kirish bloki ---
    p_tasdiq = doc.add_paragraph()
    p_tasdiq.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_tasdiq.paragraph_format.first_line_indent = Inches(0)
    p_tasdiq.paragraph_format.line_spacing = 1.15
    p_tasdiq.paragraph_format.space_after = Pt(24)
    run_tasdiq = p_tasdiq.add_run("“TASDIQLAYMAN”\n___________________________\n“____” ____________ 202__ y.")
    apply_run_font(run_tasdiq, font_name="Times New Roman", size_pt=14, bold=True)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.first_line_indent = Inches(0)
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.add_run("“AXBOROT-KUTUBXONA MARKAZI VA KUTUBXONA TARMOQLARI FAOLIYATINI BOSHQARISH VA MONITORING QILISH AXBOROT TIZIMI” (AKM-TIZIMI)")
    apply_run_font(run_title, font_name="Times New Roman", size_pt=14, bold=True)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.first_line_indent = Inches(0)
    p_sub.paragraph_format.space_before = Pt(6)
    p_sub.paragraph_format.space_after = Pt(24)
    run_sub = p_sub.add_run("DASTURIY MAJMUASINI ISHLAB CHIQISH VA JORIY ETISH BO‘YICHA\nTEXNIK TOPSHIRIQ")
    apply_run_font(run_sub, font_name="Times New Roman", size_pt=14, bold=True)

    # 1-BO'LIM
    add_heading_1("1. UMUMIY MA’LUMOTLAR")
    
    add_heading_2("1.1. Tizimning to‘liq nomi va shartli belgisi")
    add_p("Ushbu texnik topshiriqda ishlab chiqilayotgan axborot tizimining to‘liq rasmiy nomi quyidagicha belgilanadi: “Axborot-kutubxona markazi va kutubxona tarmoqlari faoliyatini boshqarish va monitoring qilish axborot tizimi”.")
    add_p("Tizimning rasmiy qisqartirilgan shartli nomi: “AKM-tizimi” (yoki “KBT — Kutubxona boshqaruv tizimi”).")

    add_heading_2("1.2. Buyurtmachi va Ishlab chiquvchi to‘g‘risida ma’lumotlar")
    add_p("Buyurtmachi: Axborot-kutubxona markazlari faoliyatini muvofiqlashtiruvchi mas’ul idora yoki viloyat axborot-kutubxona markazi ma’muriyati.")
    add_p("Ishlab chiquvchi: Dasturiy ta’minotni ishlab chiqish, sinovdan o‘tkazish va joriy etish bo‘yicha mas’ul ijrochi (dasturchilar guruhi yoki axborot texnologiyalari tashkiloti).")

    add_heading_2("1.3. Ishlarni bajarish uchun asos")
    add_p("Mazkur tizimni yaratish O‘zbekiston Respublikasining “Axborot-kutubxona faoliyati to‘g‘risida”gi Qonuni, O‘zbekiston Respublikasi Prezidentining axborot-kutubxona xizmatlari ko‘rsatishni tubdan takomillashtirishga oid qarorlari hamda davlat organlari va muassasalarida raqamlashtirishni jadallashtirish dasturlari talablari asosida amalga oshiriladi.")

    add_heading_2("1.4. Loyihani amalga oshirish bosqichlari va muddatlari")
    add_p("Loyiha doirasidagi ishlar texnik topshiriq ikki tomonlama tasdiqlangan kundan boshlab bosqichma-bosqich bajariladi. Loyihaning umumiy davomiyligi belgilangan kalendar reja-jadvalga muvofiq olib boriladi hamda tizim tahlili, loyihalash, dasturlash, integratsiya, sinovdan o‘tkazish va tajriba-sinov ekspluatatsiyasiga topshirish bosqichlarini qamrab oladi.")

    # 2-BO'LIM
    add_heading_1("2. TIZIMNING MAQSADI VA VAZIFALARI")
    
    add_heading_2("2.1. Tizimni yaratishdan ko‘zlangan asosiy maqsad")
    add_p("Axborot-kutubxona markazlari hamda ularga qarashli tuman va shahar tarmoq kutubxonalari faoliyatini to‘liq raqamlashtirish, xodimlarning ish vaqti va vazifalar ijrosi ustidan shaffof monitoring o‘rnatish, ma’lumotlarni yagona markazlashgan axborot bazasiga jamlash hamda boshqaruv qarorlarini qabul qilish samaradorligini oshirish tizimning bosh maqsadi hisoblanadi.")

    add_heading_2("2.2. Tizim orqali hal etiladigan asosiy vazifalar")
    add_p("Birinchidan, viloyat, tuman va shahar darajasidagi barcha kutubxonalar va ularning moddiy-texnik holati haqidagi ma’lumotlarni yagona reestrda shakllantirish.")
    add_p("Ikkinchidan, kutubxona xodimlarining kunlik faoliyati, xizmat ko‘rsatilgan kitobxonlar soni va bajarilgan ish turlarini avtomatlashtirilgan tarzda qayd etib borish.")
    add_p("Uchinchidan, topshiriqlar va vazifalar ijrosi intizomini mustahkamlash, ularning bajarilish muddatlarini tizimli nazorat qilish.")
    add_p("To‘rtinchidan, madaniy-ma’rifiy tadbirlarni rejalashtirish, ularning o‘tkazilishi va ishtirokchilar qamrovini monitoring qilish.")
    add_p("Beshinchidan, kitob fondi, davriy nashrlar va moddiy boyliklar inventarizatsiyasini tartibga solish.")
    add_p("Oltinchidan, xodimlar samaradorligini baholashning asosiy ko‘rsatkichlari (KPI) va reyting tizimini yo‘lga qo‘yish orqali xolis baholashni ta’minlash.")
    add_p("Yettinchidan, davriy (kunlik, oylik, choraklik va yillik) hisobotlarni qog‘ozsiz elektron shaklda yig‘ish, tekshirish va tasdiqlash mexanizmini joriy qilish.")
    add_p("Sakkizinchidan, aholi va kitobxonlardan kelib tushayotgan murojaatlar bilan ishlash jarayonini nazoratga olish.")

    add_heading_2("2.3. Kutilayotgan samaradorlik va amaliy natijalar")
    add_p("Tizimni joriy etish orqali hisobotlarni tayyorlash va umumlashtirishga ketadigan vaqt 70 foizgacha qisqaradi, inson omili bilan bog‘liq xatoliklar bartaraf etiladi, xodimlarning mehnat intizomi va topshiriqlarni o‘z vaqtida bajarish ko‘rsatkichi sezilarli darajada yaxshilanadi, viloyat va respublika miqyosida kutubxonalar faoliyati bo‘yicha tezkor tahliliy axborot olish imkoniyati yaratiladi.")

    # 3-BO'LIM
    add_heading_1("3. TIZIMGA QO‘YILADIGAN UMUMIY TALABLAR")

    add_heading_2("3.1. Tizim arxitekturasi va tuzilmasiga talablar")
    add_p("Tizim zamonaviy uch pog‘onali (mijoz, dasturiy mantiq va ma’lumotlar bazasi) arxitektura tamoyillariga asoslangan bo‘lishi lozim. Mijoz qismi (Frontend) foydalanuvchi qurilmasida veb-brauzer orqali ishlaydi. Server qismi (Backend) esa amaliy mantiqni boshqaruvchi RESTful API xizmatlarini taqdim etadi.")
    add_p("Tizim modulli tuzilishga ega bo‘lishi, kelgusida yangi funksional quyi tizimlar va modullarni mavjud arxitekturani buzmagan holda ulash imkoniyatini ta’minlashi shart.")

    add_heading_2("3.2. Dasturiy ta’minotga qo‘yiladigan talablar")
    add_p("Tizimning mijoz qismi (Frontend) zamonaviy React kutubxonasi, tezkor Vite yig‘uvchisi va Tailwind CSS uslublar tizimi asosida yasalishi shart. Foydalanuvchi interfeysi barcha zamonaviy brauzerlarda (Google Chrome, Mozilla Firefox, Microsoft Edge, Safari) bir xil to‘g‘ri va barqaror aks etishi lozim.")
    add_p("Tizimning server qismi (Backend) yuqori unumdorlikka ega Node.js va Express.js platformasida amalga oshirilishi, ma’lumotlar almashinuvi JSON formatidagi standart REST API vositasida tashkil qilinishi talab etiladi.")
    add_p("Ma’lumotlar bazasi sifatida relyatsion tuzilmali, tranzaksiyalar yaxlitligini kafolatlovchi va yuqori tezlikda ishlovchi SQLite (better-sqlite3 drayveri) yoki unga muqobil relyatsion ma’lumotlar bazasi boshqaruv tizimidan foydalaniladi.")

    add_heading_2("3.3. Texnik ta’minot va server quvvatlariga talablar")
    add_p("Server qismi uchun minimal talablar: kamida 2 yadroli protsessor, kamida 4 GB tezkor xotira (RAM), kamida 40 GB qattiq disk (SSD).")
    add_p("Foydalanuvchi ishchi o‘rni uchun talablar: internet yoki ichki tarmoqqa ulangan shaxsiy kompyuter yoki noutbuk, operatsion tizim (Windows 10/11 yoki Linux), ekranning minimal ruxsati 1280x720 piksel, yangilangan veb-brauzer.")

    add_heading_2("3.4. Ishonchlilik va barqarorlik talablari")
    add_p("Tizim sutkasiga 24 soat, haftasiga 7 kun (24/7) uzluksiz ishlash rejimiga moslashtirilishi zarur. Profilaktika va texnik xizmat ko‘rsatish ishlari foydalanuvchilar faolligi eng past bo‘lgan vaqtlarda (tungi soatlarda) amalga oshirilishi shart.")
    add_p("Tizimning nosozliklarga bardoshliligi nosozlik yuz berganda ma’lumotlarning yo‘qolmasligini ta’minlovchi ma’lumotlar bazasi tranzaksiyalari (WAL jurnali) orqali kafolatlanishi lozim.")

    add_heading_2("3.5. Foydalanuvchi interfeysi (UI/UX) va qulaylik talablari")
    add_p("Foydalanuvchi interfeysi to‘liq o‘zbek tilida (lotin alifbosida), intuitiv tushunarli, zamonaviy va ortiqcha vizual yuklamalarsiz ergonomik loyihalashtirilishi kerak.")
    add_p("Barcha asosiy amallar va funksiyalar foydalanuvchiga 2-3 marta sichqoncha bosish orqali yetib boriladigan masofada joylashtirilishi shart. Jadval ko‘rinishidagi ma’lumotlarda tezkor qidiruv, saralash, filtrlar hamda sahifalash (pagination) imkoniyatlari mavjud bo‘lishi lozim.")

    add_heading_2("3.6. Ma’lumotlar xavfsizligi va axborot himoyasi talablari")
    add_p("Tizimda axborot xavfsizligini ta’minlash maqsadida quyidagi majburiy talablar o‘rnatiladi:")
    add_p("Birinchidan, foydalanuvchilar parollari ma’lumotlar bazasida ochiq holda saqlanmaydi, ular kamida 10 raundli tuz (salt) bilan bcryptjs algoritmi orqali xeshlangan holda saqlanishi shart.")
    add_p("Ikkinchidan, foydalanuvchi sessiyalari zamonaviy JSON Web Token (JWT) mexanizmi asosida himoyalanadi va har bir so‘rovda tokenning haqiqiyligi tekshirib boriladi.")
    add_p("Uchinchidan, SQL inyeksiya, XSS (Cross-Site Scripting) va CSRF (Cross-Site Request Forgery) kabi tarmoq hujumlaridan himoyalanish uchun parametrlangan so‘rovlar va so‘rovlarni tekshiruvchi oraliq dasturlar (middleware) qo‘llaniladi.")
    add_p("To‘rtinchidan, tizimda barcha muhim harakatlar (tizimga kirish, tahrirlash, o‘chirish, tasdiqlash) Audit log jurnali orqali foydalanuvchi ID raqami, IP manzili va vaqt tamg‘asi bilan birga to‘liq qayd etib boriladi.")

    # 4-BO'LIM
    add_heading_1("4. FOYDALANUVCHILAR TARKIBI VA HUQUQLARI (RBAC TALABLARI)")

    add_heading_2("4.1. Rollar ierarxiyasi va umumiy tamoyillar")
    add_p("Tizimda foydalanuvchilarning vakolatlari Rolga asoslangan kirishni boshqarish (Role-Based Access Control — RBAC) modeli bo‘yicha taqsimlanadi. Har bir foydalanuvchiga tizimda yagona asosiy rol biriktiriladi va ularga faqat o‘z vakolati doirasidagi ma’lumotlar bilan ishlashga ruxsat etiladi.")

    add_heading_2("4.2. Super Administrator roli")
    add_p("Super Administrator tizimdagi eng yuqori vakolatga ega shaxs hisoblanadi. U tizimning barcha modullarini to‘liq boshqarish, yangi foydalanuvchilarni ro‘yxatga olish, rollar va ularning ruxsatnomalarini sozlash, tizimning umumiy parametrlarini o‘zgartirish, barcha hududiy tashkilotlar bo‘yicha to‘liq statistikani ko‘rish hamda xavfsizlik va audit jurnallarini tahlil qilish huquqiga ega.")

    add_heading_2("4.3. Viloyat Administratori roli")
    add_p("Viloyat Administratori butun viloyat miqyosidagi markaziy va tarmoq kutubxonalari faoliyatini nazorat qiladi. U o‘z viloyatiga qarashli tuman administratorlari va xodimlar boshliqlarini tayinlash, viloyat bo‘yicha umumlashtirilgan hisobotlarni tasdiqlash, tumanlar kesimidagi ko‘rsatkichlar va KPI reytinglarini tahlil qilish vakolatiga ega.")

    add_heading_2("4.4. Tuman va Shahar Administratori roli")
    add_p("Tuman va shahar administratorlari tegishli tuman yoki shahar hududida joylashgan barcha axborot-kutubxona muassasalari va filiallari faoliyatini boshqaradi. U tuman miqyosidagi xodimlarni biriktirish, ularning faoliyati bo‘yicha hisobotlarni qabul qilish hamda tuman miqyosidagi topshiriqlarni nazorat qilish huquqiga ega.")

    add_heading_2("4.5. Xodimlar Boshlig‘i roli")
    add_p("Xodimlar Boshlig‘i bevosita kutubxona xodimlarining kundalik ish faoliyatini boshqaruvchi mas’ul shaxsdir. U xodimlarga kunlik, haftalik va oylik topshiriqlarni yuklash, xodimlarning kunlik kiritgan ishlarini tekshirish va tasdiqlash, davriy hisobotlarni ko‘rib chiqib ijobiy xulosa berish yoki qayta ishlashga qaytarish, shuningdek xodimlarning KPI ko‘rsatkichlarini baholash huquqiga ega.")

    add_heading_2("4.6. Kutubxona Xodimi roli")
    add_p("Kutubxona Xodimi bevosita ijrochi hisoblanadi. U tizim orqali o‘ziga yuklatilgan topshiriqlarni ko‘rish, ijro holatini yangilash, har kuni o‘zi bajargan amaliy ishlari (kitob berish, fond bilan ishlash, elektron katalog yaratish, tadbirlarda ishtirok etish) bo‘yicha hisobot kiritish hamda o‘z shaxsiy profiling ko‘rsatkichlari bilan tanishib borish huquqiga ega.")

    # 5-BO'LIM
    add_heading_1("5. TIZIMNING FUNKSIONAL TALABLARI (QUYI TIZIMLAR VA MODULLAR)")

    add_heading_2("5.1. Autentifikatsiya va shaxsiy kabinet quyi tizimi")
    add_p("Ushbu quyi tizim foydalanuvchilarning shaxsini aniqlash va xavfsiz tizimga kirishini ta’minlaydi. Login va parol orqali tizimga kirish, noto‘g‘ri urinishlarni cheklash, faol bo‘lmagan foydalanuvchilarni bloklash, sessiya muddatini nazorat qilish hamda parolni xavfsiz yangilash funksiyalari mavjud bo‘lishi shart. Har bir foydalanuvchi o‘z shaxsiy profilida ismi-sharifi, lavozimi, biriktirilgan kutubxonasi va aloqa ma’lumotlarini ko‘rish imkoniyatiga ega bo‘ladi.")

    add_heading_2("5.2. Asosiy monitoring paneli (Dashboard) quyi tizimi")
    add_p("Dashboard quyi tizimi tizimga kirgan foydalanuvchining roliga qarab eng muhim umumlashtirilgan ko‘rsatkichlarni grafiklar va statistik kartalar ko‘rinishida aks ettiradi. Bu yerda jami kutubxonalar soni, faol xodimlar soni, kitobxonlar oqimi, kutilayotgan tadbirlar, bajarilmagan kechiktirilgan topshiriqlar hamda tasdiqlashni kutayotgan hisobotlar soni tezkor rejimda ko‘rsatiladi.")

    add_heading_2("5.3. Kutubxonalar tarmog‘ini boshqarish quyi tizimi")
    add_p("Ushbu modul barcha kutubxonalar va ularning filiallari bo‘yicha to‘liq reestrni yuritadi. Har bir kutubxona kartochkasida uning rasmiy nomi, turi, hududiy mansubligi (viloyat va tuman), manzili, geografik koordinatalari (kenglik va uzunlik), shtat birliklari soni, tashkil etilgan yili, aloqa telefonlari hamda faoliyat holati (faol, ta’mirda, yopilgan) to‘liq aks ettiriladi.")

    add_heading_2("5.4. Kitob fondi va kitobxonlar hisobi quyi tizimi")
    add_p("Kutubxonalarning adabiyotlar zaxirasi bo‘yicha to‘liq ma’lumotlar bazasi yuritiladi. Har bir adabiyotning nomi, muallifi, kategoriyasi, xalqaro standart raqami (ISBN), nashriyoti, chop etilgan yili, tili, sahifalar soni, mavjud va jami nusxalari soni qayd etiladi.")
    add_p("Kitobxonlar reestrida har bir o‘quvchining shaxsiy kartochkasi yuritiladi. Kitobxonga berilgan guvohnoma raqami, uning qarz olgan kitoblari soni, oxirgi tashrif sanasi va umumiy faolligi qayd etib boriladi.")

    add_heading_2("5.5. Xodimlarning kunlik faoliyati monitoringi quyi tizimi")
    add_p("Xodimlarning har ish kunidagi faoliyatini rejalashtirish va hisobga olish moduli bo‘lib, xodimlar tomonidan kun davomida bajarilgan ishlar (kataloglashtirish, bibliografik xizmat, kitobxonlarga xizmat ko‘rsatish, kitob ta’mirlash, ommaviy tadbirlarga tayyorgarlik) tanlangan faoliyat turi, sanasi, vaqti va batafsil tavsifi bilan tizimga kiritiladi. Xodimlar boshlig‘i ushbu qaydlarni ko‘rib chiqib, ularning haqqoniyligini tasdiqlaydi.")

    add_heading_2("5.6. Hisobotlar shakllantirish va tasdiqlash quyi tizimi")
    add_p("Quyi tizim davriy hisobotlarni shakllantirish, yuborish va tekshirish jarayonlarini to‘liq avtomatlashtiradi. Hisobotlar turi (kunlik, oylik, choraklik, yillik yoki maxsus) bo‘yicha saralanadi. Xodim hisobot loyihasini yaratadi, tegishli ko‘rsatkichlarni to‘ldiradi va rahbarga tasdiqlashga jo‘natadi. Rahbar hisobotni ma’qullashi yoki aniqlangan kamchiliklar bo‘yicha asosli izoh qoldirgan holda qayta ishlashga qaytarishi mumkin bo‘ladi.")

    add_heading_2("5.7. Topshiriqlar va ijro intizomi quyi tizimi")
    add_p("Yuqori turuvchi rahbarlar va xodimlar boshliqlari tomonidan xodimlarga aniq vazifalar yuklash moduli. Har bir topshiriq bo‘yicha uning nomi, batafsil tavsifi, ijro muddati, muhimlilik darajasi (past, o‘rta, yuqori, shoshilinch) hamda ijrochisi belgilanadi. Ijrochi topshiriq holatini (yangi, jarayonda, bajarildi) yangilab boradi.")

    add_heading_2("5.8. Madaniy-ma’rifiy tadbirlarni boshqarish quyi tizimi")
    add_p("Kutubxonalarda o‘tkaziladigan adabiy kechalar, kitob taqdimotlari, davra suhbatlari, ijodiy uchrashuvlar va tanlovlarni rejalashtirish moduli. Tadbirning nomi, yo‘nalishi, o‘tkazilish sanasi, vaqti, manzili, tashkilotchisi hamda kutilayotgan yoki amalda ishtirok etgan qatnashchilar soni to‘liq qayd etiladi.")

    add_heading_2("5.9. Moddiy-texnik baza va inventar hisobi quyi tizimi")
    add_p("Kutubxonalar tasarrufidagi asosiy vositalar va inventar buyumlari (kompyuter texnikalari, printerlar, mebel jihozlari, kitob javonlari, serverlar) hisobini yuritish moduli. Har bir inventar ob’ektining inventar raqami, sotib olingan sanasi, qiymati, soni va joriy texnik holati (yaxshi, ta’mirga muhtoj, yaroqsiz) qayd etiladi.")

    add_heading_2("5.10. Hujjatlar aylanmasi quyi tizimi")
    add_p("Kutubxona faoliyatiga oid me’yoriy hujjatlar, buyruqlar, yo‘riqnomalar va uslubiy qo‘llanmalarni tizimda markazlashgan holda saqlash va taqsimlash moduli. Hujjatlarning elektron nusxalari (PDF, DOCX va boshqa formatlarda) tegishli ruxsatnomalar asosida xodimlarga tanishish uchun taqdim etiladi.")

    add_heading_2("5.11. Fuqarolar murojaatlari bilan ishlash quyi tizimi")
    add_p("Kitobxonlar va aholidan kelib tushadigan takliflar, arizalar va shikoyatlarni qayd etish moduli. Murojaat muallifi, bog‘lanish ma’lumotlari, murojaat mazmuni, kelib tushgan sanasi, ijrochi va ijro holati (yangi, ko‘rib chiqilmoqda, hal etildi) qat’iy nazorat qilinadi.")

    add_heading_2("5.12. Xodimlar samaradorligi (KPI) va reyting quyi tizimi")
    add_p("Kutubxona xodimlarining belgilangan reja ko‘rsatkichlarini amalda bajarish darajasini baholash moduli. Har bir faoliyat yo‘nalishi bo‘yicha maqsadli ko‘rsatkich va amaldagi natija taqqoslanadi hamda foiz ko‘rsatkichlari asosida xodimlarning umumiy reytingi shakllantiriladi.")

    add_heading_2("5.13. Geolokatsiya va interaktiv xarita quyi tizimi")
    add_p("Hududiy axborot-kutubxona markazlari va tarmoq kutubxonalarining viloyat hamda tumanlar bo‘yicha geografik joylashuvini interaktiv xaritada vizual aks ettirish moduli. Xaritadagi marker orqali tegishli kutubxonaning qisqacha ma’lumotlar pasportini ko‘rish imkoniyati mavjud bo‘ladi.")

    add_heading_2("5.14. Tizim xavfsizligi, audit va loglar yuritish quyi tizimi")
    add_p("Tizimda sodir bo‘layotgan barcha jarayonlarning shaffofligini ta’minlash maqsadida har bir amal (foydalanuvchining tizimga kirishi, ma’lumot qo‘shishi, o‘zgartirishi, o‘chirishi yoki hisobotni tasdiqlashi) maxsus jurnallarda qayd etib boriladi. Ushbu yozuvlar audit tahlili uchun saqlanadi va o‘zgartirishlardan qat’iy himoyalanadi.")

    add_heading_2("5.15. Bildirishnomalar va tizim parametrlari quyi tizimi")
    add_p("Foydalanuvchilarga muhim topshiriqlar, yaqinlashib kelayotgan muddatlar, tasdiqlangan hisobotlar va tizim yangiliklari haqida real vaqt rejimida tizim ichki bildirishnomalarini yetkazish ta’minlanadi. Tizim parametrlari bo‘limi orqali esa tizim nomi, parollar uzunligi, sessiya davomiyligi kabi umumiy sozlamalar boshqariladi.")

    # 6-BO'LIM
    add_heading_1("6. MA’LUMOTLAR BAZASIGA TALABLAR")

    add_heading_2("6.1. Relyatsion tuzilma va ma’lumotlar jadvallari")
    add_p("Ma’lumotlar bazasi relyatsion arxitekturada qurilgan bo‘lib, axborotning takrorlanishini (redundancy) kamaytirish uchun uchinchi normal formaga keltiriladi. Tizim quyidagi asosiy jadvallarni o‘z ichiga oladi: foydalanuvchilar (users), kutubxonalar (libraries), kitoblar (books), kitobxonlar (readers), kunlik faoliyat (activities), hisobotlar (reports), topshiriqlar (tasks), madaniy tadbirlar (events), inventar buyumlari (inventory), hujjatlar (documents), murojaatlar (appeals), bildirishnomalar (notifications), KPI ko‘rsatkichlari (kpi_data), audit jurnali (audit_log) hamda tizim parametrlari (settings).")

    add_heading_2("6.2. Ma’lumotlar yaxlitligi va xavfsizligi talablari")
    add_p("Barcha jadvallar o‘rtasidagi bog‘liqliklar tashqi kalitlar (FOREIGN KEY) orqali nazorat qilinadi. Ma’lumotlarni yozish va o‘zgartirish operatsiyalarida relyatsion yaxlitlik ta’minlanadi. WAL (Write-Ahead Logging) rejimi orqali ma’lumotlar bazasining parallel so‘rovlarni bajarish samaradorligi va kutilmagan elektr uzilishlariga nisbatan bardoshliligi oshiriladi.")

    add_heading_2("6.3. Zaxira nusxalash (Backup) talablari")
    add_p("Ma’lumotlar bazasining to‘liq zaxira nusxasi (backup) har kuni avtomatik ravishda rejalashtirilgan jadval asosida alohida xavfsiz xotira tashuvchisiga ko‘chirilishi shart. Tizim administratoriga kerak bo‘lganda ma’lumotlarni oxirgi zaxira nusxasidan 30 daqiqa ichida qayta tiklash imkoniyati kafolatlanadi.")

    # 7-BO'LIM
    add_heading_1("7. ISHLARNI TASHKIL ETISH, SINASh VA TOPSHIRISH TARTIBI")

    add_heading_2("7.1. Ishlarni bosqichma-bosqich bajarish tartibi")
    add_p("Birinchi bosqich: Texnik topshiriqni ishlab chiqish, o‘rganish va tasdiqlash.")
    add_p("Ikkinchi bosqich: Dasturiy ta’minot arxitekturasini loyihalash, ma’lumotlar bazasi sxemasini tuzish va asosiy interfeys maketlarini yaratish.")
    add_p("Uchinchi bosqich: Server API va mijoz qismini dasturlash, barcha quyi tizim modullarini ishlab chiqish.")
    add_p("To‘rtinchi bosqich: Tizimni integratsion va funksional sinovdan o‘tkazish, yuklama ostida tekshirish hamda aniqlangan dasturiy kamchiliklarni bartaraf etish.")
    add_p("Beshinchi bosqich: Tizimni tajriba-sinov ekspluatatsiyasiga kiritish, mas’ul xodimlarni o‘qitish va foydalanish bo‘yicha qo‘llanmalarni taqdim etish.")
    add_p("Oltinchi bosqich: Yakuniy qabul qilish dalolatnomasini rasmiylashtirish va tizimni to‘liq sanoat ekspluatatsiyasiga topshirish.")

    add_heading_2("7.2. Sinovdan o‘tkazish va sifat nazorati talablari")
    add_p("Dasturiy ta’minot mustaqil sinovchi yoki mas’ul mutaxassislar tomonidan barcha funksional talablar, yuklama sinovlari, ruxsat etilmagan kirishlardan himoyalanganlik va turli brauzerlarda to‘g‘ri ishlash bo‘yicha to‘liq sinovdan o‘tkaziladi. Aniqlangan har qanday nomuvofiqliklar Ishlab chiquvchi tomonidan qisqa muddatlarda bartaraf etiladi.")

    add_heading_2("7.3. Qabul qilish va topshirish shartlari")
    add_p("Tizimni qabul qilish Buyurtmachi tomonidan tuzilgan vakolatli komissiya ishtirokida amalga oshiriladi. Tizimning mazkur Texnik topshiriqda ko‘rsatilgan barcha talablarga to‘liq javob berishi tekshirilgach, ikki tomonlama Qabul qilish-topshirish dalolatnomasi imzolanadi.")

    add_heading_2("7.4. Kafolat muddati va texnik qo‘llab-quvvatlash")
    add_p("Ishlab chiquvchi tizim qabul qilingan kundan boshlab kamida 12 (o‘n ikki) oy davomida kafolatli texnik qo‘llab-quvvatlash xizmatini taqdim etadi. Kafolat muddati davomida tizimda aniqlangan dasturiy nuqsonlar bepul bartaraf etiladi hamda tizim ma’murlariga doimiy konsultatsiya berib boriladi.")

    # 8-BO'LIM
    add_heading_1("8. LOYIHA HUJJATLARI RO‘YXATI")
    add_p("Tizimni topshirish jarayonida Ishlab chiquvchi tomonidan Buyurtmachiga quyidagi rasmiy hujjatlar to‘liq davlat tilida taqdim etiladi:")
    add_p("1. Dasturiy ta’minotning mazkur tasdiqlangan Texnik topshirig‘i;")
    add_p("2. Tizim administratorining foydalanish bo‘yicha yo‘riqnomasi;")
    add_p("3. Foydalanuvchilar (kutubxona xodimlari va rahbarlar) uchun qo‘llanma;")
    add_p("4. Dasturiy ta’minotning boshlang‘ich kodlari (source code) va o‘rnatish bo‘yicha texnik qo‘llanma;")
    add_p("5. Tizimni tajriba-sinov ekspluatatsiyasiga topshirish va yakuniy qabul qilish dalolatnomasi.")

    # Faylni saqlash
    doc.save(output_path)
    print(f"Hujjat muvaffaqiyatli saqlandi: {output_path}")

if __name__ == "__main__":
    out_file = os.path.abspath(os.path.join(r"c:\Users\abror\Desktop\AKM-tizimi-", "AKM_Tizimi_Texnik_Topshiriq.docx"))
    create_tz_document(out_file)
